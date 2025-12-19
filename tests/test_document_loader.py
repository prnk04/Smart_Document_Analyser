"""
Creating test script for Document loader
"""

import docx
import pytest
from pathlib import Path
from src.document_loader import DocumentLoader
from fpdf import FPDF


def test_document_loader_txt():
    test_file = Path("test_temp.txt")
    test_file.write_text("This is a test document about contracts.")

    doc_loader_res = DocumentLoader.load(test_file)

    assert doc_loader_res["text"] == "This is a test document about contracts."
    assert doc_loader_res["file_type"] == "TXT"
    assert doc_loader_res["word_count"] == 7

    test_file.unlink()  # Cleanup


def test_document_loader_pdf():
    # test_file = Path("test_temp.pdf")

    test_file = "test_temp.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(txt="This is a test document about contracts.", w=0)
    pdf.output(test_file)

    doc_loader_res = DocumentLoader.load(Path(test_file))

    assert doc_loader_res["text"] == "This is a test document about contracts."
    assert doc_loader_res["file_type"] == "PDF"
    assert doc_loader_res["word_count"] == 7

    Path(test_file).unlink()  # Cleanup


def test_document_loader_docx():
    document = docx.Document()
    document.add_paragraph("This is a test document about contracts.")

    # Save the changes
    document.save("test_temp.docx")

    test_file = Path("test_temp.docx")
    # test_file.write_text("This is a test document about contracts.")

    doc_loader_res = DocumentLoader.load(test_file)

    assert doc_loader_res["text"] == "This is a test document about contracts."
    assert doc_loader_res["file_type"] == "DOCX"
    assert doc_loader_res["word_count"] == 7

    test_file.unlink()  # Cleanup
